# -*- coding: utf-8 -*-
"""core.layout 的测试 —— 用自己造的矩形/文字框，不依赖任何真实单据。

线检测本身是第三方（img2table）的职责，这里只测我们自己的逻辑：
网格拼装、合并判定、文字归格、货币符号贴合、字号收敛。
"""
from core.layout import (
    CellBox, TableBlock, _build_block, blocks_to_cells, detect_blocks,
)
from core.ocr import Box
from core.writer import ExportOptions, write_workbook

XS = [50.0, 216.0, 383.0, 550.0]      # 3 列
YS = [50.0, 125.0, 200.0, 275.0, 350.0]  # 4 行


def grid_rects(rows=4, cols=3, merge_first_row=False):
    rects = []
    for r in range(rows):
        for c in range(cols):
            if merge_first_row and r == 0:
                if c == 0:
                    rects.append((XS[0], YS[0], XS[cols], YS[1]))
                continue
            rects.append((XS[c], YS[r], XS[c + 1], YS[r + 1]))
    return rects


def test_builds_grid_from_rects():
    blk = _build_block(grid_rects(), 200, [])
    assert blk is not None
    assert blk.n_rows == 4 and blk.n_cols == 3
    assert len(blk.col_widths) == 3
    assert blk.col_widths[0] == XS[1] - XS[0]


def test_merged_title_becomes_one_cell():
    """跨列的表名必须合成一个合并单元格、居中、加粗。"""
    boxes = [Box(60, 400, YS[0] + 8, YS[0] + 30, "资产 负 债 表", 0.9)]
    blk = _build_block(grid_rects(merge_first_row=True), 200, boxes)
    title = blk.cells[0]
    assert title.rowspan == 1 and title.colspan == 3
    assert title.align == "center"
    assert title.bold is True
    assert blk.n_merged >= 1


def test_text_lands_in_its_cell():
    boxes = [
        Box(60, 120, YS[1] + 8, YS[1] + 30, "货币资金", 0.9),
        Box(230, 300, YS[1] + 8, YS[1] + 30, "1,000.00", 0.9),
    ]
    blk = _build_block(grid_rects(), 200, boxes)
    by_pos = {(c.row, c.col): c for c in blk.cells}
    assert by_pos[(1, 0)].text == "货币资金"
    assert by_pos[(1, 1)].kind == "number"
    assert by_pos[(1, 1)].value == 1000.0
    assert by_pos[(1, 1)].align == "right"


def test_currency_symbol_is_glued_to_number():
    """实测：$ 常被识别成独立小框且位置略低，会拼成「101,325,990 $」解析不出来。"""
    boxes = [
        Box(220, 236, YS[1] + 12, YS[1] + 34, "$", 0.9),          # 位置略低
        Box(240, 360, YS[1] + 8, YS[1] + 30, "1,234,567", 0.9),
    ]
    blk = _build_block(grid_rects(), 200, boxes)
    cell = next(c for c in blk.cells if (c.row, c.col) == (1, 1))
    assert cell.kind == "number", cell.text
    assert cell.value == 1234567


def test_font_size_never_overflows_cell():
    """放不下时先压字号；连最小字号都放不下就标 wrap，交给 Excel 换行。"""
    long_text = "减 A deduzir： 销售成本 Custos das vendas" * 2
    boxes = [Box(60, 200, YS[1] + 4, YS[1] + 40, long_text, 0.9)]
    blk = _build_block(grid_rects(), 200, boxes)
    cell = next(c for c in blk.cells if (c.row, c.col) == (1, 0))
    assert cell.font_size >= 6.0
    assert cell.wrap is True, "长文本没标 wrap，Excel 里会溢出到隔壁格"


def test_normal_text_does_not_wrap():
    boxes = [Box(60, 200, YS[1] + 8, YS[1] + 30, "货币资金", 0.9)]
    blk = _build_block(grid_rects(), 200, boxes)
    cell = next(c for c in blk.cells if (c.row, c.col) == (1, 0))
    assert cell.wrap is False


def test_header_cell_widens_when_text_overflows():
    """表头文字比本格宽、右邻格又是空的，就该跨过去 —— 否则 Excel 里折成 3 行还挤。"""
    boxes = [Box(60, 380, YS[0] + 6, YS[0] + 30,
                 "本期 Do periodo 澳門幣（元）MOP", 0.9)]
    blk = _build_block(grid_rects(), 200, boxes)
    header = next(c for c in blk.cells if c.row == 0 and c.text)
    assert header.colspan >= 2, header.colspan


def test_last_note_row_becomes_footnote():
    long_note = ("a) 如收入中包括應徵旅遊稅或機動車輛稅之收入，須指出應稅收入之金額，"
                 "並註明應稅收入之計算依據、相關憑證編號與課稅期間，以備查核")
    boxes = [Box(60, 520, YS[3] + 6, YS[3] + 30, long_note, 0.9)]
    blk = _build_block(grid_rects(), 200, boxes)
    note = next(c for c in blk.cells if c.text and c.row == 3)
    assert note.is_footnote is True
    assert note.colspan == blk.n_cols - note.col, "脚注要跨到表格右边界"


def test_numbered_data_row_is_not_footnote():
    """数据行一律带序号（10. 銷貨成本…），即使很长也不能当脚注。"""
    long_label = ("10. 銷貨成本 Custos das vendas (1 + 2 - 3) 或 ou (4 + 5 + 6 + 7 + 8 - 9) "
                  "依本附件第三部分之規定計算")
    boxes = [Box(60, 520, YS[3] + 6, YS[3] + 30, long_label, 0.9)]
    blk = _build_block(grid_rects(), 200, boxes)
    row3 = [c for c in blk.cells if c.row == 3 and c.text]
    assert not any(c.is_footnote for c in row3), [c.text for c in row3]


def test_short_row_starting_with_marker_is_not_footnote():
    """实测坑：数据行「b) 7. 加工及分包 Subcontratos」也是 b) 开头，不能被当脚注吞掉。"""
    boxes = [Box(60, 300, YS[2] + 6, YS[2] + 30, "b) 7. 加工及分包 Subcontratos", 0.9),
             Box(230, 340, YS[2] + 8, YS[2] + 30, "$", 0.9)]
    blk = _build_block(grid_rects(), 200, boxes)
    row2 = [c for c in blk.cells if c.row == 2 and c.text]
    assert not any(c.is_footnote for c in row2), [c.text for c in row2]


def test_data_row_is_never_marked_as_footnote():
    """实测：按「最后一个数字行以下都算脚注」会把「7. 加工及分包 Subcontratos b)」吞掉。"""
    boxes = [
        Box(60, 300, YS[2] + 6, YS[2] + 30, "7. 加工及分包 Subcontratos b)", 0.9),
        Box(230, 340, YS[2] + 8, YS[2] + 30, "87,663,817", 0.9),
    ]
    blk = _build_block(grid_rects(), 200, boxes)
    row2 = [c for c in blk.cells if c.row == 2]
    assert not any(c.is_footnote for c in row2), "数据行被当成脚注了"


def test_blocks_to_cells_flattens():
    blk = TableBlock(x1=0, y1=0, x2=100, y2=100, n_rows=2, n_cols=2,
                     col_widths=[50, 50], row_heights=[50, 50],
                     cells=[CellBox(text="a", kind="text", row=0, col=0),
                            CellBox(text="b", kind="text", row=0, col=1)])
    rows = blocks_to_cells([blk])
    assert len(rows) == 2
    assert rows[0][0].text == "a"
    assert rows[1][0].kind == "empty"


def test_detect_blocks_returns_empty_when_no_table():
    """没有表格线时要老实返回空，让上层退回文本流，不能抛异常。"""
    from PIL import Image
    import io

    buf = io.BytesIO()
    Image.new("RGB", (400, 300), "white").save(buf, format="PNG")
    assert detect_blocks(buf.getvalue(), 200, ocr_boxes=[Box(1, 2, 3, 4, "x", 0.9)]) == []


def test_integer_amounts_have_no_decimals():
    """原表金额是整数，别自己加 .00 —— 多了两位会让人以为精度更高。"""
    from core.writer import _number_format_for

    assert _number_format_for(101325990.0) == "#,##0;[Red]-#,##0"
    assert _number_format_for(1234.56) == "#,##0.00;[Red]-#,##0.00"


def test_numeric_column_wide_enough_to_avoid_hashes(tmp_path):
    """实测坑：列宽按原页比例算、字号按 OCR 框估，两者不一致时
    大额数字撑破格子，Excel 显示成 ###（不是科学计数法）。"""
    big = CellBox(text="101,325,990", kind="number", value=101325990.0,
                  row=0, col=1, align="right", font_size=11)
    blk = TableBlock(x1=0, y1=0, x2=200, y2=60, n_rows=1, n_cols=2, dpi=200,
                     col_widths=[40, 40], row_heights=[40], cells=[big])

    class FakeRes:
        path = tmp_path / "x.pdf"
        page_no = 1
        source = "layout"
        cost_cny = 0.0
        review_count = 0
        blocks = [blk]
        cells = []

    out = tmp_path / "w.xlsx"
    write_workbook([FakeRes()], out, ExportOptions())

    import openpyxl
    from openpyxl.utils import get_column_letter
    wb = openpyxl.load_workbook(out)
    ws = wb[wb.sheetnames[0]]
    width = ws.column_dimensions[get_column_letter(2)].width
    assert width * 5.25 >= len("101,325,990") * 0.55 * 11, (width, "列宽不够，会显示 ###")


def test_writer_one_sheet_per_block_with_merges(tmp_path):
    merged = CellBox(text="资产负债表", kind="text", row=0, col=0,
                     colspan=3, rowspan=1, align="center", bold=True)
    body = [CellBox(text="货币资金", kind="text", row=1, col=0),
            CellBox(text="1,000.00", kind="number", value=1000.0, row=1, col=1)]
    blk = TableBlock(x1=0, y1=0, x2=300, y2=100, n_rows=2, n_cols=3, dpi=200,
                     col_widths=[100, 100, 100], row_heights=[40, 40],
                     cells=[merged] + body)

    class FakeRes:
        path = tmp_path / "x.pdf"
        page_no = 1
        source = "layout"
        cost_cny = 0.0
        review_count = 0
        blocks = [blk]
        cells = []

    out = tmp_path / "o.xlsx"
    rep = write_workbook([FakeRes()], out, ExportOptions())
    assert rep.sheets, "没有生成工作表"

    import openpyxl
    wb = openpyxl.load_workbook(out)
    ws = wb[wb.sheetnames[0]]
    assert wb.sheetnames[0] == "资产负债表", wb.sheetnames
    assert len(ws.merged_cells.ranges) == 1
    assert ws.cell(row=3, column=1).alignment.horizontal == "center"
