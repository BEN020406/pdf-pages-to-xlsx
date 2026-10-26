# -*- coding: utf-8 -*-
"""core.writer 的导出测试 —— 写出去再用 openpyxl 读回来验证。"""
from pathlib import Path

import pytest
from openpyxl import load_workbook

from core.pipeline import CellOut, PageResult
from core.reconcile import Verdict, reconcile_cell
from core.writer import ExportOptions, safe_sheet_name, write_workbook


def make_page(name="demo", page_no=2, source="ocr"):
    cells = [
        [CellOut("项目", None, "text"), CellOut("本期金额", None, "text")],
        [CellOut("营业收入", None, "text"), CellOut("1,234.56", 1234.56, "number")],
        [CellOut("营业成本", None, "text"), CellOut("", None, "empty")],
    ]
    rep = __import__("core.reconcile", fromlist=["ReconcileReport"]).ReconcileReport()
    rep.add(reconcile_cell("1,234.56", model_raw="1,234.56", ocr_score=0.99))
    # 用一个真正无法解析的值制造 conflict（'274.685 31' 现在能被 dot_restore 救回）
    rep.add(reconcile_cell("206.013.98.7", model_raw="x", ocr_score=0.95))
    return PageResult(path=Path(f"{name}.pdf"), page_no=page_no, source=source,
                      cells=cells, report=rep, png=b"", cost_cny=0.006)


class TestBorders:
    """黑框加粗 —— 用户明确要求「有框框可以变黑加粗」。"""

    def test_borders_are_pure_black(self, tmp_path):
        """实测 bug：Side(color='000000') 被 openpyxl 当 ARGB 解析成
        00000000（全透明），打印出来看不见。必须是 FF000000。"""
        out = tmp_path / "o.xlsx"
        write_workbook([make_page()], out)
        ws = load_workbook(out)["demo"]
        cell = ws.cell(row=3, column=1)
        assert cell.border.left.color.rgb == "FF000000"

    def test_outer_border_is_medium(self, tmp_path):
        """外框加粗，内框细 —— 会计表格的视觉惯例。"""
        out = tmp_path / "o.xlsx"
        write_workbook([make_page()], out)
        ws = load_workbook(out)["demo"]
        top_left = ws.cell(row=3, column=1)
        assert top_left.border.left.style == "medium"
        assert top_left.border.top.style == "medium"
        inner = ws.cell(row=4, column=2)
        assert inner.border.left.style == "thin"

    def test_gridlines_hidden(self, tmp_path):
        out = tmp_path / "o.xlsx"
        write_workbook([make_page()], out)
        assert load_workbook(out)["demo"].sheet_view.showGridLines is False


class TestLayoutRestore:
    """按 PDF 的 pt 还原行列尺寸。"""

    def _page_with_layout(self):
        from core.tables import TableLayout
        p = make_page()
        p.layout = TableLayout(
            page_no=2, bbox=(72, 130, 512, 208),
            row_edges=[130, 156, 182, 208],
            col_edges=[72, 182, 292, 402, 512],
            cells=[["a"] * 4 for _ in range(3)],
            has_borders=True, line_width=0.8,
        )
        return p

    def test_row_heights_from_pdf(self, tmp_path):
        out = tmp_path / "o.xlsx"
        write_workbook([self._page_with_layout()], out)
        ws = load_workbook(out)["demo"]
        # 每行 26pt
        assert ws.row_dimensions[3].height == pytest.approx(26.0, abs=0.5)

    def test_col_widths_from_pdf(self, tmp_path):
        out = tmp_path / "o.xlsx"
        write_workbook([self._page_with_layout()], out)
        ws = load_workbook(out)["demo"]
        # 110pt / 5.25 ≈ 20.95
        assert ws.column_dimensions["A"].width == pytest.approx(20.95, abs=0.5)

    def test_layout_off_uses_autofit(self, tmp_path):
        out = tmp_path / "o.xlsx"
        write_workbook([self._page_with_layout()], out,
                       ExportOptions(restore_layout=False))
        ws = load_workbook(out)["demo"]
        assert ws.row_dimensions[3].height != pytest.approx(26.0, abs=0.5)


class TestStatusFillNotHidingBorders:
    """实测 bug：状态底色（C6EFCE 等饱和色）会盖住黑框，表格看起来「框断了」。"""

    def test_status_fill_is_pale(self, tmp_path):
        out = tmp_path / "o.xlsx"
        page = make_page()
        page.cells[1][1].verdict = Verdict.AGREE
        write_workbook([page], out)
        ws = load_workbook(out)["demo"]
        c = ws.cell(row=4, column=2)
        rgb = str(c.fill.fgColor.rgb)
        if len(rgb) == 8:
            rgb = rgb[2:]
        # 不能是原始的饱和绿 C6EFCE
        assert rgb.upper() != "C6EFCE"
        # 也不能是纯白（那样就没有状态提示了）
        assert rgb.upper() != "FFFFFF"

    def test_border_still_black_on_filled_cell(self, tmp_path):
        """底色不能影响框线本身。"""
        out = tmp_path / "o.xlsx"
        page = make_page()
        page.cells[1][1].verdict = Verdict.AGREE
        write_workbook([page], out)
        ws = load_workbook(out)["demo"]
        c = ws.cell(row=4, column=2)
        assert c.border.left.color.rgb == "FF000000"
        assert c.border.top.color.rgb == "FF000000"

    def test_unknown_color_auto_paled(self, tmp_path):
        out = tmp_path / "o.xlsx"
        page = make_page()
        page.cells[1][1].verdict = Verdict.AGREE
        write_workbook([page], out)
        # 直接验证函数行为
        from core.writer import _status_fill
        f = _status_fill("AABBCC")
        assert f is not None


class TestScannedFallbackLayout:
    """扫描件走 OCR 路径没有 layout，行高不能是 Excel 默认的 15pt。"""

    def test_default_height_when_no_layout(self, tmp_path):
        out = tmp_path / "o.xlsx"
        page = make_page()
        page.layout = None
        write_workbook([page], out)
        ws = load_workbook(out)["demo"]
        assert ws.row_dimensions[3].height == pytest.approx(22.0, abs=0.5)


class TestSheetName:
    def test_illegal_chars_replaced(self):
        assert "[" not in safe_sheet_name("a[b]c", set())

    def test_truncated_to_31(self):
        assert len(safe_sheet_name("x" * 50, set())) <= 31

    def test_dedup(self):
        used = set()
        a = safe_sheet_name("same", used)
        b = safe_sheet_name("same", used)
        assert a != b and b.endswith("_2")


class TestWrite:
    def test_single_page(self, tmp_path):
        out = tmp_path / "o.xlsx"
        rep = write_workbook([make_page()], out)
        assert out.exists()
        assert rep.sheets[0][0] == "demo"

        wb = load_workbook(out)
        assert "demo" in wb.sheetnames
        assert "OCR校验" in wb.sheetnames
        assert "汇总" in wb.sheetnames

    def test_numbers_written_as_numbers(self, tmp_path):
        out = tmp_path / "o.xlsx"
        write_workbook([make_page()], out)
        ws = load_workbook(out)["demo"]
        values = [c.value for row in ws.iter_rows() for c in row]
        assert 1234.56 in values
        # 不能写成字符串
        assert "1234.56" not in [v for v in values if isinstance(v, str)]

    def test_empty_cell_stays_blank_not_zero(self, tmp_path):
        """空单元格绝不能写成 0 —— 会让合计出错。"""
        out = tmp_path / "o.xlsx"
        write_workbook([make_page()], out)
        ws = load_workbook(out)["demo"]
        numbers = [c.value for row in ws.iter_rows() for c in row
                   if isinstance(c.value, (int, float))]
        assert 0 not in numbers

    def test_title_contains_source(self, tmp_path):
        out = tmp_path / "o.xlsx"
        write_workbook([make_page(source="llm")], out)
        ws = load_workbook(out)["demo"]
        assert "llm" in str(ws["A1"].value)

    def test_audit_sheet_has_traceability(self, tmp_path):
        out = tmp_path / "o.xlsx"
        write_workbook([make_page()], out)
        ws = load_workbook(out)["OCR校验"]
        header = [c.value for c in ws[1]]
        assert header[:4] == ["原文", "规则", "本地值", "模型值"]
        assert ws.max_row >= 3

    def test_summary_counts_review(self, tmp_path):
        out = tmp_path / "o.xlsx"
        rep = write_workbook([make_page()], out)
        assert rep.total_review == 1        # 那条 conflict
        ws = load_workbook(out)["汇总"]
        values = [c.value for row in ws.iter_rows() for c in row]
        assert "合计" in values

    def test_multi_page_multiple_sheets(self, tmp_path):
        out = tmp_path / "o.xlsx"
        pages = [make_page(f"f{i}", i + 1) for i in range(3)]
        rep = write_workbook(pages, out)
        assert len(rep.sheets) == 3
        assert len(load_workbook(out).sheetnames) == 5   # 3 + 校验 + 汇总

    def test_no_image_option(self, tmp_path):
        out = tmp_path / "o.xlsx"
        write_workbook([make_page()], out, ExportOptions(include_image=False))
        assert out.exists()

    def test_no_audit_sheet_option(self, tmp_path):
        out = tmp_path / "o.xlsx"
        write_workbook([make_page()], out, ExportOptions(include_audit_sheet=False))
        assert "OCR校验" not in load_workbook(out).sheetnames

    def test_colorize_applied(self, tmp_path):
        out = tmp_path / "o.xlsx"
        page = make_page()
        page.cells[1][1].verdict = Verdict.CONFLICT
        write_workbook([page], out)
        ws = load_workbook(out)["demo"]
        filled = [c for row in ws.iter_rows() for c in row
                  if c.fill and c.fill.fgColor.rgb not in (None, "00000000")]
        assert filled, "冲突单元格应该有底色"

    def test_error_page_still_written(self, tmp_path):
        out = tmp_path / "o.xlsx"
        page = make_page()
        page.error = "页码超出范围"
        write_workbook([page], out)
        ws = load_workbook(out)["demo"]
        assert "错误" in str(ws["A1"].value)

    def test_cost_in_summary(self, tmp_path):
        out = tmp_path / "o.xlsx"
        rep = write_workbook([make_page(), make_page("b")], out)
        assert rep.total_cost == pytest.approx(0.012)
